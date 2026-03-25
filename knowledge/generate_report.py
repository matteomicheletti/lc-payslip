#!/usr/bin/env python3
"""
Generate Excel report of overtime 80/20 split bug discrepancies.
"""

import csv
import math
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

CSV_PATH = "/home/matteo/Documents/dev/lc-payslip/data.csv"
OUTPUT_PATH = "/home/matteo/Documents/dev/lc-payslip/knowledge/report_discrepanze_straordinario.xlsx"

# --- Parse CSV and compute discrepancies (same logic as find_discrepancies.py) ---

with open(CSV_PATH, encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    rows = list(reader)

emp_month_day = defaultdict(lambda: defaultdict(list))

for row in rows:
    giorno_inizio = (row.get("GIORNO INIZIO") or "").strip()
    nome = (row.get("NOME DIPENDENTE") or "").strip()
    if not giorno_inizio or not nome:
        continue
    parts = giorno_inizio.split("-")
    if len(parts) != 3:
        continue
    dd, mm, yyyy = parts
    month_key = f"{mm}-{yyyy}"
    emp_month_day[(nome, month_key)][giorno_inizio].append(row)


def parse_float(val):
    if not val:
        return 0.0
    val = val.strip().replace(",", ".")
    try:
        return float(val)
    except ValueError:
        return 0.0


results = []

for (employee, month), days_dict in sorted(emp_month_day.items()):
    total_ord_min = 0.0
    total_straord_min = 0.0
    pos = None

    for day, day_rows in days_dict.items():
        day_ord = 0.0
        day_straord = 0.0
        for r in day_rows:
            day_ord += parse_float(r.get("MIN. ORD. VAL"))
            day_straord += parse_float(r.get("MIN. STRAORD. VAL"))
            if pos is None:
                p = parse_float(r.get("POS"))
                if p > 0:
                    pos = p
        if day_ord > 480:
            excess = day_ord - 480
            day_ord = 480
            day_straord += excess
        total_ord_min += day_ord
        total_straord_min += day_straord

    if pos is None:
        pos = 0.0

    total_straord_hours = total_straord_min / 60.0

    if total_straord_hours <= 5:
        continue

    full_importo_straord = total_straord_hours * pos

    # OLD
    ore_straord_old = round(total_straord_hours * 0.8)
    ore_IB_old = round(total_straord_hours * 0.2)
    importo_straord_old = round(full_importo_straord * 0.8)
    importo_IB_old = round(full_importo_straord * 0.2)

    # NEW
    ore_straord_new = round(total_straord_hours * 0.8)
    ore_IB_new = round(total_straord_hours) - ore_straord_new
    importo_straord_new = ore_straord_new * pos
    importo_IB_new = ore_IB_new * pos

    old_total = importo_straord_old + importo_IB_old
    new_total = importo_straord_new + importo_IB_new
    delta = new_total - old_total

    has_diff = (
        ore_straord_old != ore_straord_new
        or ore_IB_old != ore_IB_new
        or abs(importo_straord_old - importo_straord_new) > 0.001
        or abs(importo_IB_old - importo_IB_new) > 0.001
    )

    if has_diff:
        # Extract year from month key for sorting
        mm_str, yyyy_str = month.split("-")
        results.append({
            "employee": employee,
            "month": month,
            "mm": int(mm_str),
            "yyyy": int(yyyy_str),
            "total_straord_hours": total_straord_hours,
            "pos": pos,
            "full_importo": full_importo_straord,
            "ore_straord_old": ore_straord_old,
            "ore_IB_old": ore_IB_old,
            "importo_straord_old": importo_straord_old,
            "importo_IB_old": importo_IB_old,
            "old_total": old_total,
            "ore_straord_new": ore_straord_new,
            "ore_IB_new": ore_IB_new,
            "importo_straord_new": importo_straord_new,
            "importo_IB_new": importo_IB_new,
            "new_total": new_total,
            "delta": delta,
        })

# Sort by employee, then year, then month
results.sort(key=lambda r: (r["employee"], r["yyyy"], r["mm"]))

# --- Build Excel ---

wb = Workbook()

# ==================== SHEET 1: Detail ====================
ws = wb.active
ws.title = "Dettaglio Discrepanze"

# Styles
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill_old = PatternFill(start_color="C0392B", end_color="C0392B", fill_type="solid")
header_fill_new = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
header_fill_info = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
header_fill_delta = PatternFill(start_color="8E44AD", end_color="8E44AD", fill_type="solid")
positive_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
negative_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

headers = [
    ("Dipendente", header_fill_info),
    ("Mese", header_fill_info),
    ("Anno", header_fill_info),
    ("Ore Straord. Totali", header_fill_info),
    ("POS (EUR/h)", header_fill_info),
    ("Importo Straord. Pieno", header_fill_info),
    ("Ore Straord. (80%) OLD", header_fill_old),
    ("Ore IB (20%) OLD", header_fill_old),
    ("PLUS Straord. OLD", header_fill_old),
    ("Importo IB OLD", header_fill_old),
    ("Totale OLD", header_fill_old),
    ("Ore Straord. (80%) NEW", header_fill_new),
    ("Ore IB (20%) NEW", header_fill_new),
    ("PLUS Straord. NEW", header_fill_new),
    ("Importo IB NEW", header_fill_new),
    ("Totale NEW", header_fill_new),
    ("DELTA (EUR)", header_fill_delta),
]

ITALIAN_MONTHS = {
    1: "Gennaio", 2: "Febbraio", 3: "Marzo", 4: "Aprile",
    5: "Maggio", 6: "Giugno", 7: "Luglio", 8: "Agosto",
    9: "Settembre", 10: "Ottobre", 11: "Novembre", 12: "Dicembre",
}

# Write headers
for col_idx, (header_text, fill) in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header_text)
    cell.font = header_font
    cell.fill = fill
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = thin_border

# Write data rows
eur_fmt = '#,##0.00'
for row_idx, r in enumerate(results, 2):
    values = [
        r["employee"],
        ITALIAN_MONTHS[r["mm"]],
        r["yyyy"],
        round(r["total_straord_hours"], 2),
        r["pos"],
        round(r["full_importo"], 2),
        r["ore_straord_old"],
        r["ore_IB_old"],
        r["importo_straord_old"],
        r["importo_IB_old"],
        r["old_total"],
        r["ore_straord_new"],
        r["ore_IB_new"],
        round(r["importo_straord_new"], 2),
        round(r["importo_IB_new"], 2),
        round(r["new_total"], 2),
        round(r["delta"], 2),
    ]
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")
        # Format EUR columns
        if col_idx in (5, 6, 9, 10, 11, 14, 15, 16, 17):
            cell.number_format = eur_fmt
        # Color delta column
        if col_idx == 17:
            if val > 0:
                cell.fill = positive_fill
                cell.font = Font(bold=True, color="27AE60")
            elif val < 0:
                cell.fill = negative_fill
                cell.font = Font(bold=True, color="C0392B")

# Auto-fit column widths
for col_idx in range(1, len(headers) + 1):
    max_len = len(str(headers[col_idx - 1][0]))
    for row_idx in range(2, len(results) + 2):
        cell_val = ws.cell(row=row_idx, column=col_idx).value
        if cell_val is not None:
            max_len = max(max_len, len(str(cell_val)))
    ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 25)

# Freeze header row
ws.freeze_panes = "A2"
# Auto-filter
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"

# ==================== SHEET 2: Summary per employee ====================
ws2 = wb.create_sheet("Riepilogo per Dipendente")

emp_summary = defaultdict(lambda: {"count": 0, "total_delta": 0.0, "months": []})
for r in results:
    key = r["employee"]
    emp_summary[key]["count"] += 1
    emp_summary[key]["total_delta"] += r["delta"]
    emp_summary[key]["months"].append(f"{ITALIAN_MONTHS[r['mm']]} {r['yyyy']}")

summary_headers = [
    "Dipendente",
    "Mesi Interessati",
    "Delta Totale (EUR)",
    "Dettaglio Mesi",
]

for col_idx, h in enumerate(summary_headers, 1):
    cell = ws2.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill_info
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = thin_border

for row_idx, (emp, data) in enumerate(sorted(emp_summary.items()), 2):
    values = [
        emp,
        data["count"],
        round(data["total_delta"], 2),
        ", ".join(data["months"]),
    ]
    for col_idx, val in enumerate(values, 1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=val)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if col_idx != 4 else "left", wrap_text=(col_idx == 4))
        if col_idx == 3:
            cell.number_format = eur_fmt
            if val > 0:
                cell.fill = positive_fill
                cell.font = Font(bold=True, color="27AE60")
            elif val < 0:
                cell.fill = negative_fill
                cell.font = Font(bold=True, color="C0392B")

ws2.column_dimensions["A"].width = 25
ws2.column_dimensions["B"].width = 18
ws2.column_dimensions["C"].width = 20
ws2.column_dimensions["D"].width = 60
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = f"A1:D{len(emp_summary) + 1}"

# ==================== SHEET 3: Grand total ====================
ws3 = wb.create_sheet("Totale Generale")

total_positive = sum(r["delta"] for r in results if r["delta"] > 0)
total_negative = sum(r["delta"] for r in results if r["delta"] < 0)
grand_total = sum(r["delta"] for r in results)

summary_rows = [
    ("Casi con discrepanza trovati", len(results), None),
    ("Dipendenti interessati", len(emp_summary), None),
    ("", "", None),
    ("Totale sottopagato (dipendente ha ricevuto meno del dovuto)", round(total_positive, 2), positive_fill),
    ("Totale sovrapagato (dipendente ha ricevuto piu del dovuto)", round(abs(total_negative), 2), negative_fill),
    ("", "", None),
    ("SALDO NETTO (positivo = sottopagato, negativo = sovrapagato)", round(grand_total, 2), None),
]

title_font = Font(bold=True, size=14, color="2C3E50")
cell = ws3.cell(row=1, column=1, value="REPORT DISCREPANZE — CALCOLO STRAORDINARIO 80/20")
cell.font = title_font

subtitle_font = Font(italic=True, size=10, color="7F8C8D")
cell = ws3.cell(row=2, column=1, value="Differenza tra vecchio calcolo (bug) e nuovo calcolo (corretto)")
cell.font = subtitle_font

for row_idx, (label, value, fill) in enumerate(summary_rows, 4):
    label_cell = ws3.cell(row=row_idx, column=1, value=label)
    label_cell.font = Font(bold=True, size=11)
    label_cell.border = thin_border

    val_cell = ws3.cell(row=row_idx, column=2, value=value)
    val_cell.border = thin_border
    val_cell.alignment = Alignment(horizontal="center")
    if isinstance(value, float):
        val_cell.number_format = eur_fmt
        val_cell.font = Font(bold=True, size=11)
    if fill:
        val_cell.fill = fill

# Grand total row special styling
grand_cell_label = ws3.cell(row=10, column=1)
grand_cell_label.font = Font(bold=True, size=13, color="8E44AD")
grand_cell_val = ws3.cell(row=10, column=2)
grand_cell_val.font = Font(bold=True, size=13, color="8E44AD")
if grand_total > 0:
    grand_cell_val.fill = positive_fill
elif grand_total < 0:
    grand_cell_val.fill = negative_fill

ws3.column_dimensions["A"].width = 60
ws3.column_dimensions["B"].width = 20

# Save
wb.save(OUTPUT_PATH)
print(f"Report saved to: {OUTPUT_PATH}")
print(f"  - {len(results)} discrepancies across {len(emp_summary)} employees")
print(f"  - Grand total delta: {grand_total:.2f} EUR")
print(f"    (sottopagato: +{total_positive:.2f} | sovrapagato: {total_negative:.2f})")
